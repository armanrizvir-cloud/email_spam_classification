from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

def apply_theme(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(24, 33, 45) # Dark Navy Blue

def style_title(title_shape):
    if not title_shape.has_text_frame: return
    for p in title_shape.text_frame.paragraphs:
        for run in p.runs:
            run.font.size = Pt(40)
            run.font.name = 'Segoe UI'
            run.font.bold = True
            run.font.color.rgb = RGBColor(64, 224, 208) # Turquoise / Cyan

def style_subtitle(subtitle_shape):
    if not subtitle_shape.has_text_frame: return
    for p in subtitle_shape.text_frame.paragraphs:
        for run in p.runs:
            run.font.size = Pt(22)
            run.font.name = 'Segoe UI'
            run.font.color.rgb = RGBColor(220, 220, 220)

def create_presentation():
    prs = Presentation()
    
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    apply_theme(slide)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "MailGuard: Intelligent Email Spam Classification"
    subtitle.text = "A Machine Learning approach using NLP\n\nPresented by: Arman Rizvi\n3rd Year B.Tech, Data Science & AI"
    
    style_title(title)
    style_subtitle(subtitle)

    def add_slide(title_text, bullet_points):
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        apply_theme(slide)
        
        title = slide.shapes.title
        title.text = title_text
        style_title(title)
        
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        
        for point in bullet_points:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0
            for run in p.runs:
                run.font.size = Pt(22)
                run.font.name = 'Segoe UI'
                run.font.color.rgb = RGBColor(240, 240, 240)
            
    add_slide("Problem Statement", [
        "Email spam wastes time, reduces productivity, and often carries malicious payloads.",
        "Traditional rule-based filters are easily bypassed by modern spammers.",
        "There is a need for an automated, NLP-based detection tool that understands context."
    ])

    add_slide("Project Objectives", [
        "Preprocess and extract features from unstructured email text.",
        "Train and compare multiple Machine Learning classification algorithms.",
        "Evaluate models using Precision and F1-score to minimize false positives.",
        "Develop an automated pipeline for intelligent model selection.",
        "Build an interactive web dashboard for real-time predictions."
    ])

    add_slide("Dataset", [
        "Sourced from HuggingFace / Kaggle (10,900+ emails).",
        "Features used: Raw email text.",
        "Target Variable: 'label' (0 = Ham/Legitimate, 1 = Spam).",
        "Class Distribution: Perfectly balanced (~50% Spam, 50% Ham)."
    ])

    add_slide("Data Preprocessing & NLP", [
        "Handling Missing Data: Filled empty text fields.",
        "NLP Cleaning: Converted to lowercase, removing formatting.",
        "Vectorization: Used TF-IDF (Term Frequency-Inverse Document Frequency).",
        "N-grams: Extracted both unigrams and bigrams for context."
    ])

    add_slide("Algorithms Evaluated", [
        "1. Logistic Regression: A fast, interpretable baseline.",
        "2. Naive Bayes: The classic NLP text classifier.",
        "3. Random Forest: An ensemble method to capture complex patterns.",
        "4. Linear SVC: A Support Vector Machine optimized for sparse TF-IDF matrices."
    ])

    add_slide("Evaluation Metrics", [
        "Accuracy is not enough. We prioritized:",
        "Precision: Out of all flagged emails, how many were actually spam? (Crucial to avoid filtering legitimate emails)",
        "Recall: Out of all actual spam, how many did we catch?",
        "F1-Score: The harmonic mean of Precision and Recall. Our primary metric for model selection."
    ])

    add_slide("Training Results", [
        "Logistic Regression: Accuracy 99.6% | Precision 99.8%",
        "Naive Bayes: Accuracy 99.5% | Precision 99.3%",
        "Random Forest: Accuracy 99.3% | Precision 98.9%",
        "Linear SVC: Accuracy 99.6% | Precision 99.7%",
        "Winner: Linear SVC. Automatically selected by the pipeline for balancing top-tier Precision and F1-Score."
    ])

    add_slide("Interactive Dashboard", [
        "Built using Python Dash and Dash Bootstrap Components.",
        "Features a 'Cyborg' dark-mode premium UI.",
        "Real-time Inference: Users paste an email, and the backend vectorizes and predicts instantly.",
        "Model Insights: Visualizes exactly which words the model flagged as 'Spam'."
    ])

    add_slide("Conclusion & Future Scope", [
        "Conclusion:",
        "- Successfully built an end-to-end ML pipeline.",
        "- Proved that Linear SVC is highly effective on sparse text data.",
        "- Achieved 99.7% Precision, protecting legitimate emails.",
        "Future Scope:",
        "- Upgrade from TF-IDF to Transformer models (BERT).",
        "- Deploy as a web-mail API plugin."
    ])

    os.makedirs('docs', exist_ok=True)
    save_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'docs', 'MailGuard_Presentation.pptx')
    prs.save(save_path)
    print(f"Presentation generated successfully at: {save_path}")

if __name__ == "__main__":
    create_presentation()
